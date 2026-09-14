#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
=============================================================
  ANALIZADOR DE COBROS — MERCADO PAGO
  Configurá los parámetros en la sección "CONFIGURACIÓN"
=============================================================

Toma el export de cobros de Mercado Pago y devuelve un Excel con los cobros
agrupados. La particularidad del rubro gastronómico está contemplada:
un cobro de las 2 o 3 de la mañana pertenece al TURNO DEL DÍA ANTERIOR.

USO
---
  python analizador_cobros_mp.py collection-2024.xls
  python analizador_cobros_mp.py enero.xls febrero.xls marzo.xls   # varios meses
  python analizador_cobros_mp.py                # toma el export más nuevo de la carpeta
  python analizador_cobros_mp.py archivo.xls --diagnostico   # solo analiza, no escribe

En Colab:
  !python analizador_cobros_mp.py "/content/collection-2024.xls"

SALIDA (7 hojas)
----------------
  1. Cobros por Día      -> día de turno, con la regla de la madrugada
  2. Cobros por Hora     -> ordenado según el turno, no según el reloj
  3. Resumen Mensual     -> bruto, neto, retenciones, ticket promedio
  4. Medios de Pago      -> participación de cada medio por mes
  5. Tarifas e Impuestos -> qué se descontó y cuánto pesa sobre el bruto
  6. No Concretadas      -> rechazos y cancelaciones, por motivo
  7. Detalle Cobros      -> operación por operación, con su día de turno

Los resúmenes usan SUMIFS / COUNTIFS contra "Detalle Cobros": si filtrás o
corregís algo ahí, los totales se recalculan solos.
"""

import argparse
import glob
import os
import re
import sys
import unicodedata
import xml.etree.ElementTree as ET
from datetime import datetime, timedelta

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# =============================================================================
#  CONFIGURACIÓN  <- editá esto sin tocar el resto
# =============================================================================

ARCHIVO_ENTRADA = None      # None = toma el export más nuevo de la carpeta
ARCHIVO_SALIDA = None       # None = analisis_cobros_<períodos>.xlsx

# --- LA REGLA DEL RUBRO ---------------------------------------------------
# Todo cobro anterior a esta hora se imputa al turno del día anterior.
# Con 6: un cobro de las 02:30 del sábado pertenece al turno del viernes.
# En los datos de prueba no hay un solo cobro entre las 3 y las 7 de la
# mañana, así que cualquier valor entre 3 y 7 da el mismo resultado.
# Poné 7 si querés replicar exactamente la ventana 7:00 a 2:00.
HORA_CORTE = 6

# --- QUÉ CUENTA COMO COBRO ------------------------------------------------
# Se clasifica POR EXCLUSIÓN, a propósito: si Mercado Pago suma un tipo de
# operación nuevo (por ejemplo cobros con lector Point además del QR), entra
# solo como venta en vez de quedar afuera sin que nadie lo note. El script
# avisa por consola cuando aparece un tipo que no conocía.
ESTADOS_COBRO = ["approved"]
OPERACIONES_QUE_NO_SON_VENTA = [
    "account_fund",        # carga de saldo a tu propia cuenta
    "money_transfer",      # transferencia enviada
    "withdrawal",          # retiro a cuenta bancaria
    "payout",
    "money_exchange",
    "credit_payment",      # pago de un crédito de MP
]
OPERACIONES_VENTA_CONOCIDAS = [
    "regular_payment",     # QR, link de pago, checkout
    "point_payment",       # lector físico Point
    "pos_payment",
    "subscription_payment",
]

ETIQUETAS_MEDIO_PAGO = {
    "account_money": "Dinero en cuenta (saldo MP)",
    "bank_transfer": "Transferencia / débito inmediato",
    "credit_card": "Tarjeta de crédito",
    "debit_card": "Tarjeta de débito",
    "prepaid_card": "Tarjeta prepaga",
    "digital_currency": "Moneda digital / cripto",
    "ticket": "Pago en efectivo (cupón)",
    "atm": "Cajero automático",
}

ETIQUETAS_MOTIVO = {
    "cc_rejected_call_for_authorize": "Rechazo: requiere autorización del banco",
    "cc_rejected_card_disabled": "Rechazo: tarjeta deshabilitada",
    "cc_rejected_other_reason": "Rechazo: motivo no especificado por el banco",
    "cc_rejected_insufficient_amount": "Rechazo: fondos insuficientes",
    "cc_rejected_time_out": "Rechazo: tiempo de espera agotado",
    "cc_rejected_bad_filled_security_code": "Rechazo: código de seguridad incorrecto",
    "cc_rejected_bad_filled_date": "Rechazo: fecha de tarjeta incorrecta",
    "cc_rejected_high_risk": "Rechazo: riesgo detectado",
    "cc_rejected_bad_filled_card_number": "Rechazo: número de tarjeta incorrecto",
    "expired": "Cancelada: QR / cobro expirado sin pago",
    "cancelled": "Cancelada por el vendedor o el comprador",
}

DIAS_SEMANA = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]

# Tarifas explícitas que trae el archivo.
CLAVES_TARIFAS = ["mercadopago_fee", "marketplace_fee", "shipping_cost", "financing_fee"]

# =============================================================================
#  ESTILOS
# =============================================================================
FUENTE = "Arial"
F_TITULO = Font(name=FUENTE, size=14, bold=True, color="00437A")
F_SUBTIT = Font(name=FUENTE, size=9, italic=True, color="595959")
F_CABECERA = Font(name=FUENTE, size=10, bold=True, color="FFFFFF")
F_NORMAL = Font(name=FUENTE, size=10)
F_TOTAL = Font(name=FUENTE, size=10, bold=True)
FILL_CAB = PatternFill("solid", fgColor="00437A")
FILL_TOTAL = PatternFill("solid", fgColor="D6E4F0")
FILL_ALT = PatternFill("solid", fgColor="F2F2F2")
FILL_FINDE = PatternFill("solid", fgColor="FFF2CC")
BORDE_TOP = Border(top=Side(style="thin", color="808080"))

FMT_PESOS = '"$" #,##0.00;[RED]-"$" #,##0.00;"-"'
FMT_PCT = '0.0%;-0.0%;"-"'
FMT_FECHA = "DD/MM/YYYY"
FMT_HORA = "HH:MM"
FMT_FECHAHORA = "DD/MM/YYYY HH:MM"
FMT_ENT = '#,##0;-#,##0;"-"'


# =============================================================================
#  LECTURA
# =============================================================================
def _leer_spreadsheetml(path: str) -> pd.DataFrame:
    """El '.xls' de Mercado Pago suele ser XML de Excel 2003, no un Excel real.

    Respeta ss:Index: las celdas vacías no se escriben en el XML, se saltean
    con un índice. Si no se contempla, las columnas se corren y los importes
    terminan en la columna equivocada.
    """
    ns = {"ss": "urn:schemas-microsoft-com:office:spreadsheet"}
    attr_index = "{urn:schemas-microsoft-com:office:spreadsheet}Index"
    filas = []
    for _, el in ET.iterparse(path, events=("end",)):
        if el.tag.split("}")[-1] != "Row":
            continue
        fila, pos = {}, 0
        for celda in el:
            if celda.tag.split("}")[-1] != "Cell":
                continue
            if attr_index in celda.attrib:
                pos = int(celda.attrib[attr_index]) - 1
            dato = celda.find("ss:Data", ns)
            fila[pos] = dato.text if dato is not None else None
            pos += 1
        filas.append(fila)
        el.clear()
    if not filas:
        raise ValueError("el XML no tiene filas")
    ancho = max((max(f) for f in filas if f), default=-1) + 1
    matriz = [[f.get(i) for i in range(ancho)] for f in filas]
    return pd.DataFrame(matriz[1:], columns=matriz[0])


def leer_export(path: str) -> pd.DataFrame:
    """Abre el export probando los formatos en los que Mercado Pago lo entrega."""
    errores = []

    # ¿es XML de Excel 2003 aunque diga .xls?
    try:
        with open(path, "rb") as fh:
            arranque = fh.read(200).lstrip()
        if arranque.startswith(b"<?xml") or arranque.startswith(b"<Workbook"):
            return _leer_spreadsheetml(path)
    except Exception as e:                                       # noqa: BLE001
        errores.append(f"spreadsheetml: {e}")

    for motor in ("openpyxl", "calamine", "xlrd"):
        try:
            return pd.read_excel(path, engine=motor)
        except Exception as e:                                   # noqa: BLE001
            errores.append(f"{motor}: {type(e).__name__}")
    try:
        return pd.read_csv(path, sep=None, engine="python", encoding="latin-1")
    except Exception as e:                                       # noqa: BLE001
        errores.append(f"csv: {e}")

    raise SystemExit(f"[ERROR] No pude abrir {os.path.basename(path)}.\n"
                     "        " + "\n        ".join(errores) +
                     "\n        Si falta 'calamine': pip install python-calamine")


def renombrar_a_clave_interna(df: pd.DataFrame) -> pd.DataFrame:
    """'Medio de pago (payment_type)' -> 'payment_type'.

    La clave interna es más estable que el texto en español, que cambia de
    redacción entre versiones del reporte.
    """
    nuevos = {}
    for col in df.columns:
        m = re.search(r"\(([^()]+)\)\s*$", str(col).strip())
        nuevos[col] = m.group(1).strip() if m else str(col).strip()
    return df.rename(columns=nuevos)


def a_numero(valor) -> float:
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return 0.0
    if isinstance(valor, (int, float)):
        return float(valor)
    t = str(valor).strip()
    if not t:
        return 0.0
    negativo = t.startswith("(") and t.endswith(")")
    t = re.sub(r"[^\d,.\-]", "", t.strip("()"))
    if "," in t and "." in t:
        t = t.replace(".", "").replace(",", ".")
    elif "," in t:
        t = t.replace(",", ".")
    try:
        n = float(t)
    except ValueError:
        return 0.0
    return -n if negativo else n


def columna_segura(df: pd.DataFrame, clave: str) -> pd.Series:
    """La columna convertida a número, o ceros si este export no la trae."""
    if clave in df.columns:
        return df[clave].map(a_numero)
    return pd.Series(0.0, index=df.index)


# =============================================================================
#  LA REGLA DEL TURNO
# =============================================================================
def dia_de_turno(momento: pd.Timestamp, hora_corte: int = HORA_CORTE):
    """Devuelve a qué turno pertenece un cobro.

    Un cobro de las 02:30 del sábado es del turno del viernes.
    """
    if pd.isna(momento):
        return pd.NaT
    fecha = momento.normalize()
    return fecha - timedelta(days=1) if momento.hour < hora_corte else fecha


def orden_horas_del_turno(hora_corte: int = HORA_CORTE):
    """Las 24 horas arrancando por la apertura: 6,7,...,23,0,1,...,5."""
    return [(hora_corte + i) % 24 for i in range(24)]


# =============================================================================
#  ARMADO DEL DATASET
# =============================================================================
def cargar(rutas) -> tuple:
    partes = []
    for ruta in rutas:
        crudo = renombrar_a_clave_interna(leer_export(ruta))
        crudo["archivo_origen"] = os.path.basename(ruta)
        partes.append(crudo)
        print(f"    {os.path.basename(ruta)}: {len(crudo)} filas")
    datos = pd.concat(partes, ignore_index=True, sort=False)

    # varios exports pueden solaparse: una operación se cuenta una sola vez
    if "operation_id" in datos.columns:
        antes = len(datos)
        datos["operation_id"] = datos["operation_id"].astype(str)
        datos = datos.drop_duplicates(subset="operation_id", keep="first")
        if antes - len(datos):
            print(f"[i] Descarto {antes - len(datos)} operaciones duplicadas entre archivos.")

    # --- fecha y hora ---
    col_fecha = next((c for c in ("date_created", "date_approved", "date_created_short")
                      if c in datos.columns), None)
    if col_fecha is None:
        raise SystemExit("[ERROR] El archivo no trae ninguna columna de fecha reconocible.")
    datos["momento"] = pd.to_datetime(datos[col_fecha], dayfirst=True, errors="coerce")
    datos = datos[datos["momento"].notna()]

    tiene_hora = bool((datos["momento"].dt.time != pd.Timestamp("00:00:00").time()).any())
    if not tiene_hora:
        print("[!] ATENCIÓN: el archivo trae la fecha SIN hora, así que no se puede\n"
              "    aplicar el corte de turno. Cada cobro queda en su día calendario.\n"
              "    Descargá el export con 'Fecha de compra (date_created)' completa.")

    datos["dia_turno"] = (datos["momento"].map(lambda m: dia_de_turno(m, HORA_CORTE))
                          if tiene_hora else datos["momento"].dt.normalize())
    datos["hora"] = datos["momento"].dt.hour
    datos["dia_semana"] = datos["dia_turno"].dt.weekday.map(lambda i: DIAS_SEMANA[i])
    datos["periodo"] = datos["dia_turno"].dt.to_period("M").astype(str)

    # --- qué es un cobro ---
    estado = datos.get("status", pd.Series("approved", index=datos.index)).astype(str)
    tipo = datos.get("operation_type", pd.Series("regular_payment", index=datos.index)).astype(str)

    desconocidos = sorted(set(tipo.unique())
                          - set(OPERACIONES_VENTA_CONOCIDAS)
                          - set(OPERACIONES_QUE_NO_SON_VENTA))
    if desconocidos:
        print(f"[!] Tipos de operación que no tenía clasificados: {desconocidos}\n"
              f"    Los cuento COMO VENTA. Si alguno no lo es, agregalo a "
              f"OPERACIONES_QUE_NO_SON_VENTA.")

    es_cobro = estado.isin(ESTADOS_COBRO) & ~tipo.isin(OPERACIONES_QUE_NO_SON_VENTA)
    es_no_concretada = estado.isin(["rejected", "cancelled"])
    es_fondeo = tipo.isin(OPERACIONES_QUE_NO_SON_VENTA) & estado.isin(ESTADOS_COBRO)

    cobros = datos[es_cobro].copy()
    cobros["bruto"] = columna_segura(cobros, "transaction_amount")
    cobros["neto"] = columna_segura(cobros, "net_received_amount")
    cobros["comision_mp"] = columna_segura(cobros, "mercadopago_fee").abs()
    cobros["otras_tarifas"] = sum(columna_segura(cobros, c).abs()
                                  for c in CLAVES_TARIFAS if c != "mercadopago_fee")
    # Lo que no está discriminado en ninguna columna: bruto - tarifas - neto.
    # En Argentina suele ser percepciones/retenciones (IIBB) que MP aplica como
    # agente de recaudación pero no desglosa. Conviene validarlo con contaduría.
    cobros["retenciones"] = (cobros["bruto"] - cobros["comision_mp"]
                             - cobros["otras_tarifas"] - cobros["neto"]).round(2)
    cobros["medio_pago"] = (cobros.get("payment_type", pd.Series("", index=cobros.index))
                            .astype(str).map(ETIQUETAS_MEDIO_PAGO)
                            .fillna(cobros.get("payment_type", pd.Series("sin dato",
                                                                         index=cobros.index))
                                    .astype(str).str.replace("_", " ").str.title()))
    cobros["tipo_operacion"] = tipo[es_cobro]
    cobros["nro_operacion"] = datos.get("operation_id", pd.Series("", index=datos.index))[es_cobro]

    no_concretadas = datos[es_no_concretada].copy()
    if len(no_concretadas):
        no_concretadas["bruto"] = columna_segura(no_concretadas, "transaction_amount")
        detalle = no_concretadas.get("status_detail",
                                     pd.Series("", index=no_concretadas.index)).astype(str)
        no_concretadas["motivo"] = detalle.map(ETIQUETAS_MOTIVO).fillna(detalle)

    fondeos = datos[es_fondeo].copy()
    if len(fondeos):
        fondeos["bruto"] = columna_segura(fondeos, "transaction_amount")

    cobros = cobros.sort_values("momento").reset_index(drop=True)
    return cobros, no_concretadas, fondeos, tiene_hora


# =============================================================================
#  ESCRITURA
# =============================================================================
def encabezado_hoja(ws, titulo, subtitulo, ancho):
    ws["A1"] = titulo
    ws["A1"].font = F_TITULO
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ancho)
    ws["A2"] = subtitulo
    ws["A2"].font = F_SUBTIT
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ancho)
    ws.row_dimensions[1].height = 20


def fila_cabecera(ws, fila, cabeceras):
    for j, h in enumerate(cabeceras, start=1):
        c = ws.cell(row=fila, column=j, value=h)
        c.font, c.fill = F_CABECERA, FILL_CAB
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def escribir_detalle(wb, cobros):
    """Hoja base. Las fórmulas de todas las demás dependen de estas columnas:
       A momento | B día turno | C día semana | D hora | E período
       F medio pago | G tipo op | H bruto | I comisión | J otras | K retenc | L neto"""
    ws = wb.create_sheet("Detalle Cobros")
    cabeceras = ["Fecha y hora", "Día de turno", "Día", "Hora", "Período", "Medio de pago",
                 "Tipo de operación", "Bruto", "Comisión MP", "Otras tarifas",
                 "Retenciones (est.)", "Neto recibido", "Nº operación"]
    fila_cabecera(ws, 1, cabeceras)

    for _, r in cobros.iterrows():
        ws.append([r["momento"].to_pydatetime(), r["dia_turno"].to_pydatetime(),
                   r["dia_semana"], int(r["hora"]), r["periodo"], r["medio_pago"],
                   r["tipo_operacion"], r["bruto"], r["comision_mp"], r["otras_tarifas"],
                   r["retenciones"], r["neto"], str(r["nro_operacion"])])

    ultima = ws.max_row
    for f in range(2, ultima + 1):
        for c in range(1, len(cabeceras) + 1):
            ws.cell(row=f, column=c).font = F_NORMAL
        ws.cell(row=f, column=1).number_format = FMT_FECHAHORA
        ws.cell(row=f, column=2).number_format = FMT_FECHA
        for c in (8, 9, 10, 11, 12):
            ws.cell(row=f, column=c).number_format = FMT_PESOS

    for col, ancho in zip("ABCDEFGHIJKLM",
                          [18, 13, 11, 7, 10, 30, 18, 15, 14, 14, 16, 15, 16]):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cabeceras))}{ultima}"
    return ultima


def escribir_cobros_por_dia(wb, cobros, fin, subtitulo):
    """La hoja del pedido: cobros por día de turno."""
    ws = wb.create_sheet("Cobros por Día")
    cabeceras = ["Día de turno", "Día", "Cobros", "Bruto", "Neto recibido",
                 "Ticket promedio", "% s/total", "Primer cobro", "Último cobro"]
    encabezado_hoja(ws, "Cobros por día de turno",
                    subtitulo + f"  |  Corte de turno: {HORA_CORTE:02d}:00 "
                    f"(los cobros de 00:00 a {HORA_CORTE - 1:02d}:59 van al día anterior)",
                    len(cabeceras))
    fc = 4
    fila_cabecera(ws, fc, cabeceras)

    r_dia = f"'Detalle Cobros'!$B$2:$B${fin}"
    r_bruto = f"'Detalle Cobros'!$H$2:$H${fin}"
    r_neto = f"'Detalle Cobros'!$L$2:$L${fin}"

    por_dia = (cobros.groupby("dia_turno")
               .agg(dia_semana=("dia_semana", "first"),
                    primero=("momento", "min"), ultimo=("momento", "max"))
               .reset_index().sort_values("dia_turno"))

    f = fc + 1
    primera = f
    for _, r in por_dia.iterrows():
        ws.cell(row=f, column=1, value=r["dia_turno"].to_pydatetime())
        ws.cell(row=f, column=2, value=r["dia_semana"])
        ws.cell(row=f, column=3, value=f"=COUNTIFS({r_dia},$A{f})")
        ws.cell(row=f, column=4, value=f"=SUMIFS({r_bruto},{r_dia},$A{f})")
        ws.cell(row=f, column=5, value=f"=SUMIFS({r_neto},{r_dia},$A{f})")
        ws.cell(row=f, column=6, value=f'=IFERROR(D{f}/C{f},"")')
        ws.cell(row=f, column=8, value=r["primero"].to_pydatetime())
        ws.cell(row=f, column=9, value=r["ultimo"].to_pydatetime())
        f += 1
    ultima, total = f - 1, f

    ws.cell(row=total, column=1, value="TOTAL")
    for col in ("C", "D", "E"):
        ws[f"{col}{total}"] = f"=SUM({col}{primera}:{col}{ultima})"
    ws[f"F{total}"] = f'=IFERROR(D{total}/C{total},"")'
    for r in range(primera, ultima + 1):
        ws.cell(row=r, column=7, value=f'=IFERROR(D{r}/$D${total},"")')
    ws.cell(row=total, column=7, value=f'=IFERROR(SUM(G{primera}:G{ultima}),"")')

    for r in range(primera, total + 1):
        es_total = r == total
        finde = ws.cell(row=r, column=2).value in ("Sábado", "Domingo")
        for c in range(1, len(cabeceras) + 1):
            celda = ws.cell(row=r, column=c)
            celda.font = F_TOTAL if es_total else F_NORMAL
            if es_total:
                celda.fill, celda.border = FILL_TOTAL, BORDE_TOP
            elif finde:
                celda.fill = FILL_FINDE
        ws.cell(row=r, column=1).number_format = FMT_FECHA
        ws.cell(row=r, column=3).number_format = FMT_ENT
        for c in (4, 5, 6):
            ws.cell(row=r, column=c).number_format = FMT_PESOS
        ws.cell(row=r, column=7).number_format = FMT_PCT
        for c in (8, 9):
            ws.cell(row=r, column=c).number_format = FMT_HORA

    for col, ancho in zip("ABCDEFGHI", [14, 12, 10, 17, 17, 16, 11, 12, 12]):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = f"A{fc + 1}"
    ws.auto_filter.ref = f"A{fc}:I{ultima}"


def escribir_cobros_por_hora(wb, cobros, fin, subtitulo):
    """Las 24 horas, pero ordenadas según el turno y no según el reloj."""
    ws = wb.create_sheet("Cobros por Hora")
    cabeceras = ["Hora", "Cobros", "Bruto", "% s/total", "Ticket promedio"]
    encabezado_hoja(ws, "Cobros por hora del turno",
                    subtitulo + f"  |  El listado arranca a las {HORA_CORTE:02d}:00; "
                    "las horas del final son la madrugada del día siguiente",
                    len(cabeceras))
    fc = 4
    fila_cabecera(ws, fc, cabeceras)

    r_hora = f"'Detalle Cobros'!$D$2:$D${fin}"
    r_bruto = f"'Detalle Cobros'!$H$2:$H${fin}"
    presentes = set(cobros["hora"].unique())

    f = fc + 1
    primera = f
    for hora in orden_horas_del_turno(HORA_CORTE):
        if hora not in presentes:
            continue
        etiqueta = f"{hora:02d}:00 a {hora:02d}:59"
        if hora < HORA_CORTE:
            etiqueta += "  (madrugada)"
        ws.cell(row=f, column=1, value=etiqueta)
        ws.cell(row=f, column=2, value=f"=COUNTIFS({r_hora},{hora})")
        ws.cell(row=f, column=3, value=f"=SUMIFS({r_bruto},{r_hora},{hora})")
        ws.cell(row=f, column=5, value=f'=IFERROR(C{f}/B{f},"")')
        f += 1
    ultima, total = f - 1, f

    ws.cell(row=total, column=1, value="TOTAL")
    for col in ("B", "C"):
        ws[f"{col}{total}"] = f"=SUM({col}{primera}:{col}{ultima})"
    ws[f"E{total}"] = f'=IFERROR(C{total}/B{total},"")'
    for r in range(primera, ultima + 1):
        ws.cell(row=r, column=4, value=f'=IFERROR(C{r}/$C${total},"")')
    ws.cell(row=total, column=4, value=f'=IFERROR(SUM(D{primera}:D{ultima}),"")')

    for r in range(primera, total + 1):
        es_total = r == total
        madrugada = "madrugada" in str(ws.cell(row=r, column=1).value)
        for c in range(1, len(cabeceras) + 1):
            celda = ws.cell(row=r, column=c)
            celda.font = F_TOTAL if es_total else F_NORMAL
            if es_total:
                celda.fill, celda.border = FILL_TOTAL, BORDE_TOP
            elif madrugada:
                celda.fill = FILL_FINDE
        ws.cell(row=r, column=2).number_format = FMT_ENT
        ws.cell(row=r, column=3).number_format = FMT_PESOS
        ws.cell(row=r, column=4).number_format = FMT_PCT
        ws.cell(row=r, column=5).number_format = FMT_PESOS

    for col, ancho in zip("ABCDE", [24, 10, 18, 11, 16]):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = f"A{fc + 1}"


def escribir_resumen_mensual(wb, cobros, fin, subtitulo):
    ws = wb.create_sheet("Resumen Mensual")
    cabeceras = ["Período", "Cobros", "Bruto", "Comisión MP", "Otras tarifas",
                 "Retenciones (est.)", "Neto recibido", "Total descontado",
                 "% descontado", "Ticket promedio"]
    encabezado_hoja(ws, "Resumen mensual", subtitulo, len(cabeceras))
    fc = 4
    fila_cabecera(ws, fc, cabeceras)

    r_per = f"'Detalle Cobros'!$E$2:$E${fin}"
    cols = {"H": 3, "I": 4, "J": 5, "K": 6, "L": 7}

    f = fc + 1
    primera = f
    for periodo in sorted(cobros["periodo"].unique()):
        ws.cell(row=f, column=1, value=periodo)
        ws.cell(row=f, column=2, value=f"=COUNTIFS({r_per},$A{f})")
        for letra, col in cols.items():
            rango = f"'Detalle Cobros'!${letra}$2:${letra}${fin}"
            ws.cell(row=f, column=col, value=f"=SUMIFS({rango},{r_per},$A{f})")
        ws.cell(row=f, column=8, value=f"=C{f}-G{f}")
        ws.cell(row=f, column=9, value=f'=IFERROR(H{f}/C{f},"")')
        ws.cell(row=f, column=10, value=f'=IFERROR(C{f}/B{f},"")')
        f += 1
    ultima, total = f - 1, f

    ws.cell(row=total, column=1, value="TOTAL")
    for col in ("B", "C", "D", "E", "F", "G", "H"):
        ws[f"{col}{total}"] = f"=SUM({col}{primera}:{col}{ultima})"
    ws[f"I{total}"] = f'=IFERROR(H{total}/C{total},"")'
    ws[f"J{total}"] = f'=IFERROR(C{total}/B{total},"")'

    for r in range(primera, total + 1):
        es_total = r == total
        for c in range(1, len(cabeceras) + 1):
            celda = ws.cell(row=r, column=c)
            celda.font = F_TOTAL if es_total else F_NORMAL
            if es_total:
                celda.fill, celda.border = FILL_TOTAL, BORDE_TOP
        ws.cell(row=r, column=2).number_format = FMT_ENT
        for c in (3, 4, 5, 6, 7, 8, 10):
            ws.cell(row=r, column=c).number_format = FMT_PESOS
        ws.cell(row=r, column=9).number_format = FMT_PCT

    for col, ancho in zip("ABCDEFGHIJ", [12, 10, 18, 16, 15, 17, 18, 17, 13, 16]):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = f"A{fc + 1}"


def escribir_medios_de_pago(wb, cobros, fin, subtitulo):
    ws = wb.create_sheet("Medios de Pago")
    cabeceras = ["Período", "Medio de pago", "Cobros", "Bruto", "% del período",
                 "Ticket promedio"]
    encabezado_hoja(ws, "Participación por medio de pago", subtitulo, len(cabeceras))
    fc = 4
    fila_cabecera(ws, fc, cabeceras)

    r_per = f"'Detalle Cobros'!$E$2:$E${fin}"
    r_medio = f"'Detalle Cobros'!$F$2:$F${fin}"
    r_bruto = f"'Detalle Cobros'!$H$2:$H${fin}"

    grupos = (cobros.groupby(["periodo", "medio_pago"])["bruto"].sum()
              .reset_index().sort_values(["periodo", "bruto"], ascending=[True, False]))

    f = fc + 1
    primera = f
    filas_por_periodo = {}
    for _, r in grupos.iterrows():
        ws.cell(row=f, column=1, value=r["periodo"])
        ws.cell(row=f, column=2, value=r["medio_pago"])
        ws.cell(row=f, column=3, value=f"=COUNTIFS({r_per},$A{f},{r_medio},$B{f})")
        ws.cell(row=f, column=4, value=f"=SUMIFS({r_bruto},{r_per},$A{f},{r_medio},$B{f})")
        ws.cell(row=f, column=5,
                value=f'=IFERROR(D{f}/SUMIFS({r_bruto},{r_per},$A{f}),"")')
        ws.cell(row=f, column=6, value=f'=IFERROR(D{f}/C{f},"")')
        filas_por_periodo.setdefault(r["periodo"], []).append(f)
        f += 1
    ultima, total = f - 1, f

    ws.cell(row=total, column=1, value="TOTAL")
    for col in ("C", "D"):
        ws[f"{col}{total}"] = f"=SUM({col}{primera}:{col}{ultima})"
    ws[f"F{total}"] = f'=IFERROR(D{total}/C{total},"")'

    for r in range(primera, total + 1):
        es_total = r == total
        for c in range(1, len(cabeceras) + 1):
            celda = ws.cell(row=r, column=c)
            celda.font = F_TOTAL if es_total else F_NORMAL
            if es_total:
                celda.fill, celda.border = FILL_TOTAL, BORDE_TOP
            elif (r - primera) % 2 == 1:
                celda.fill = FILL_ALT
        ws.cell(row=r, column=3).number_format = FMT_ENT
        ws.cell(row=r, column=4).number_format = FMT_PESOS
        ws.cell(row=r, column=5).number_format = FMT_PCT
        ws.cell(row=r, column=6).number_format = FMT_PESOS

    for col, ancho in zip("ABCDEF", [12, 34, 10, 18, 14, 16]):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = f"A{fc + 1}"


def escribir_tarifas(wb, cobros, fin, subtitulo):
    ws = wb.create_sheet("Tarifas e Impuestos")
    cabeceras = ["Período", "Concepto", "Monto", "% s/bruto"]
    encabezado_hoja(
        ws, "Tarifas, comisiones y retenciones",
        subtitulo + "  |  'Retenciones no discriminadas' es la diferencia entre "
        "bruto − tarifas explícitas − neto acreditado: conviene validarla con contaduría",
        len(cabeceras))
    fc = 4
    fila_cabecera(ws, fc, cabeceras)

    r_per = f"'Detalle Cobros'!$E$2:$E${fin}"
    r_bruto = f"'Detalle Cobros'!$H$2:$H${fin}"
    conceptos = [("Comisión Mercado Pago", "I"),
                 ("Otras tarifas (envío, plataforma, financiación)", "J"),
                 ("Retenciones no discriminadas (estimado)", "K")]

    f = fc + 1
    primera = f
    for periodo in sorted(cobros["periodo"].unique()):
        for etiqueta, letra in conceptos:
            rango = f"'Detalle Cobros'!${letra}$2:${letra}${fin}"
            ws.cell(row=f, column=1, value=periodo)
            ws.cell(row=f, column=2, value=etiqueta)
            ws.cell(row=f, column=3, value=f"=SUMIFS({rango},{r_per},$A{f})")
            ws.cell(row=f, column=4,
                    value=f'=IFERROR(C{f}/SUMIFS({r_bruto},{r_per},$A{f}),"")')
            f += 1
    ultima, total = f - 1, f

    ws.cell(row=total, column=1, value="TOTAL")
    ws[f"C{total}"] = f"=SUM(C{primera}:C{ultima})"
    ws[f"D{total}"] = f'=IFERROR(C{total}/SUM({r_bruto}),"")'

    for r in range(primera, total + 1):
        es_total = r == total
        for c in range(1, len(cabeceras) + 1):
            celda = ws.cell(row=r, column=c)
            celda.font = F_TOTAL if es_total else F_NORMAL
            if es_total:
                celda.fill, celda.border = FILL_TOTAL, BORDE_TOP
        ws.cell(row=r, column=3).number_format = FMT_PESOS
        ws.cell(row=r, column=4).number_format = FMT_PCT

    for col, ancho in zip("ABCD", [12, 48, 18, 13]):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = f"A{fc + 1}"


def escribir_no_concretadas(wb, no_concretadas, subtitulo):
    ws = wb.create_sheet("No Concretadas")
    cabeceras = ["Período", "Motivo", "Cantidad", "Monto no cobrado"]
    encabezado_hoja(ws, "Cobros rechazados y cancelados",
                    subtitulo + "  |  No son facturación, pero sirven como alerta operativa",
                    len(cabeceras))
    fc = 4
    fila_cabecera(ws, fc, cabeceras)

    resumen = (no_concretadas.groupby(["periodo", "motivo"])
               .agg(cantidad=("bruto", "size"), monto=("bruto", "sum"))
               .reset_index().sort_values(["periodo", "monto"], ascending=[True, False]))

    f = fc + 1
    primera = f
    for _, r in resumen.iterrows():
        ws.cell(row=f, column=1, value=r["periodo"])
        ws.cell(row=f, column=2, value=r["motivo"])
        ws.cell(row=f, column=3, value=int(r["cantidad"]))
        ws.cell(row=f, column=4, value=float(r["monto"]))
        f += 1
    ultima, total = f - 1, f

    ws.cell(row=total, column=1, value="TOTAL")
    for col in ("C", "D"):
        ws[f"{col}{total}"] = f"=SUM({col}{primera}:{col}{ultima})"

    for r in range(primera, total + 1):
        es_total = r == total
        for c in range(1, len(cabeceras) + 1):
            celda = ws.cell(row=r, column=c)
            celda.font = F_TOTAL if es_total else F_NORMAL
            if es_total:
                celda.fill, celda.border = FILL_TOTAL, BORDE_TOP
        ws.cell(row=r, column=3).number_format = FMT_ENT
        ws.cell(row=r, column=4).number_format = FMT_PESOS

    for col, ancho in zip("ABCD", [12, 46, 12, 20]):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = f"A{fc + 1}"


# =============================================================================
#  INFORME POR CONSOLA
# =============================================================================
def informe(cobros, no_concretadas, fondeos, tiene_hora):
    print("\n" + "=" * 74)
    print("  RESUMEN")
    print("=" * 74)
    print(f"  Cobros                  : {len(cobros):>8,}   "
          f"$ {cobros['bruto'].sum():>16,.2f}")
    print(f"  Neto acreditado         : {'':>8}   $ {cobros['neto'].sum():>16,.2f}")
    print(f"  Descontado por MP       : {'':>8}   "
          f"$ {cobros['bruto'].sum() - cobros['neto'].sum():>16,.2f}  "
          f"({(1 - cobros['neto'].sum() / cobros['bruto'].sum()) * 100:.2f}%)")
    if len(no_concretadas):
        print(f"  No concretadas          : {len(no_concretadas):>8,}   "
              f"$ {no_concretadas['bruto'].sum():>16,.2f}  (no es facturación)")
    if len(fondeos):
        print(f"  Cargas de cuenta        : {len(fondeos):>8,}   "
              f"$ {fondeos['bruto'].sum():>16,.2f}  (no son ventas)")

    if tiene_hora:
        madrugada = cobros[cobros["hora"] < HORA_CORTE]
        print(f"\n  REGLA DEL TURNO (corte {HORA_CORTE:02d}:00)")
        if len(madrugada):
            pct = madrugada["bruto"].sum() / cobros["bruto"].sum() * 100
            print(f"    {len(madrugada):,} cobros por $ {madrugada['bruto'].sum():,.2f} "
                  f"({pct:.1f}% del total) se movieron al día anterior.")
            por_hora = madrugada.groupby("hora")["bruto"].agg(["size", "sum"])
            for hora, r in por_hora.iterrows():
                print(f"      {hora:02d}h: {int(r['size']):>4} cobros  "
                      f"$ {r['sum']:>14,.2f}")
        else:
            print("    No hubo cobros de madrugada en este período.")

        horas_sin = [h for h in range(24) if h not in set(cobros["hora"].unique())]
        if horas_sin:
            tramos = []
            inicio = anterior = horas_sin[0]
            for h in horas_sin[1:]:
                if h == anterior + 1:
                    anterior = h
                else:
                    tramos.append((inicio, anterior))
                    inicio = anterior = h
            tramos.append((inicio, anterior))
            texto = ", ".join(f"{a:02d}h" if a == b else f"{a:02d}h a {b:02d}h"
                              for a, b in tramos)
            print(f"    Sin actividad: {texto}  "
                  f"(sirve para confirmar dónde conviene el corte)")

    mejor = (cobros.groupby(["dia_turno", "dia_semana"])["bruto"].sum()
             .sort_values(ascending=False))
    if len(mejor):
        (fecha, dia), monto = mejor.index[0], mejor.iloc[0]
        print(f"\n  Mejor turno: {dia} {fecha:%d/%m/%Y} con $ {monto:,.2f}")
        promedio = cobros.groupby("dia_semana")["bruto"].sum() / \
            cobros.groupby("dia_semana")["dia_turno"].nunique()
        print("  Promedio por día de la semana:")
        for dia in DIAS_SEMANA:
            if dia in promedio.index:
                print(f"    {dia.ljust(11)} $ {promedio[dia]:>14,.2f}")
    print("=" * 74 + "\n")


# =============================================================================
#  FLUJO PRINCIPAL
# =============================================================================
def generar(rutas, path_salida=None, solo_diagnostico=False):
    print("Leyendo:")
    cobros, no_concretadas, fondeos, tiene_hora = cargar(rutas)
    if cobros.empty:
        raise SystemExit("[ERROR] No se encontró ningún cobro aprobado en el archivo.")

    informe(cobros, no_concretadas, fondeos, tiene_hora)
    if solo_diagnostico:
        return None

    periodos = sorted(cobros["periodo"].unique())
    if path_salida is None:
        path_salida = f"analisis_cobros_{'-'.join(periodos)}.xlsx"

    f_min, f_max = cobros["dia_turno"].min(), cobros["dia_turno"].max()
    subtitulo = (f"Turnos del {f_min:%d/%m/%Y} al {f_max:%d/%m/%Y}  |  "
                 f"{len(cobros):,} cobros  |  "
                 f"Generado: {datetime.now():%d/%m/%Y %H:%M}")

    wb = Workbook()
    wb.remove(wb.active)
    fin = escribir_detalle(wb, cobros)          # primero: todo lo demás la referencia
    escribir_cobros_por_dia(wb, cobros, fin, subtitulo)
    escribir_cobros_por_hora(wb, cobros, fin, subtitulo)
    escribir_resumen_mensual(wb, cobros, fin, subtitulo)
    escribir_medios_de_pago(wb, cobros, fin, subtitulo)
    escribir_tarifas(wb, cobros, fin, subtitulo)
    if len(no_concretadas):
        escribir_no_concretadas(wb, no_concretadas, subtitulo)

    orden = ["Cobros por Día", "Cobros por Hora", "Resumen Mensual", "Medios de Pago",
             "Tarifas e Impuestos", "No Concretadas", "Detalle Cobros"]
    wb._sheets = [wb[n] for n in orden if n in wb.sheetnames]
    wb.active = 0
    wb.save(path_salida)

    print(f"--> {path_salida}")
    return path_salida


def main():
    ap = argparse.ArgumentParser(description="Analiza los cobros de Mercado Pago por turno.")
    ap.add_argument("entradas", nargs="*", help="uno o más exports de cobros")
    ap.add_argument("-o", "--salida", default=ARCHIVO_SALIDA)
    ap.add_argument("--hora-corte", type=int, default=None,
                    help=f"hora de corte del turno (default {HORA_CORTE})")
    ap.add_argument("--diagnostico", action="store_true",
                    help="solo analiza, no escribe el Excel")
    args = ap.parse_args()

    if args.hora_corte is not None:
        if not 0 <= args.hora_corte <= 23:
            sys.exit("[ERROR] La hora de corte tiene que estar entre 0 y 23.")
        globals()["HORA_CORTE"] = args.hora_corte

    rutas = args.entradas or ([ARCHIVO_ENTRADA] if ARCHIVO_ENTRADA else [])
    if not rutas:
        candidatos = [f for f in glob.glob("*.xls") + glob.glob("*.xlsx")
                      if not f.startswith("~$") and "analisis_cobros" not in f.lower()]
        if not candidatos:
            sys.exit("[ERROR] No hay ningún export en esta carpeta. Pasá la ruta como argumento.")
        rutas = [max(candidatos, key=os.path.getmtime)]
        print(f"[i] Uso el archivo más reciente: {rutas[0]}")

    faltantes = [r for r in rutas if not os.path.exists(r)]
    if faltantes:
        sys.exit(f"[ERROR] No existe: {', '.join(faltantes)}")

    generar(rutas, args.salida, solo_diagnostico=args.diagnostico)


if __name__ == "__main__":
    main()
